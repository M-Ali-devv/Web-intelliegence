"""Phase 4: search, review, and streamed exports."""

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import get_session
from app.main import app
from app.models import Website


@pytest.fixture()
def client():
    return TestClient(app)


def test_search_filters_by_company_name(make_website, client):
    make_website("https://acme.test", status="classified", company_name="ACME Robotics", industry="Manufacturing / Industrial")
    make_website("https://other.test", status="classified", company_name="Other Co", industry="Education")

    response = client.get("/", params={"q": "ACME"})

    assert "https://acme.test" in response.text
    assert "https://other.test" not in response.text
    assert "Find companies" in response.text


def test_blank_confidence_does_not_break_search(make_website, client):
    make_website("https://acme.test", status="classified", company_name="ACME Robotics")

    response = client.get("/", params={"q": "ACME", "min_confidence": ""})

    assert response.status_code == 200
    assert "https://acme.test" in response.text


def test_review_queue_lists_low_confidence_rows(make_website, client):
    make_website("https://sure.test", status="classified", company_name="Sure Co", confidence=90)
    make_website(
        "https://unsure.test",
        status="needs_review",
        company_name="Unsure Co",
        confidence=60,
        evidence=json.dumps([{"field": "industry", "quote": "We build robots.", "page": "about"}]),
    )

    response = client.get("/", params={"status": "review"})

    assert "https://unsure.test" in response.text
    assert "https://sure.test" not in response.text
    assert "We build robots." in response.text
    assert "Accept" in response.text


def test_accept_and_reject_review(make_website, client):
    site_id = make_website("https://unsure.test", status="needs_review", company_name="Unsure Co", confidence=60)

    accepted = client.post(f"/review/{site_id}", data={"decision": "accept"}, follow_redirects=True)
    assert "Review saved" in accepted.text
    with get_session() as session:
        site = session.get(Website, site_id)
        assert site.status == "classified"
        assert json.loads(site.classification_meta)["review_decision"] == "accepted"

    client.post(f"/review/{site_id}", data={"decision": "reject"}, follow_redirects=False)
    with get_session() as session:
        site = session.get(Website, site_id)
        assert site.status == "insufficient_evidence"
        assert json.loads(site.classification_meta)["review_decision"] == "rejected"


def test_retry_classification_uses_the_classifier(monkeypatch, make_website, client):
    site_id = make_website("https://acme.test", status="classify_failed", clean_text="ACME builds robots for factories worldwide.")
    called = {}

    def fake_classify(site, provider=None):
        called["id"] = site.id
        site.status = "classified"
        site.company_name = "ACME"

    monkeypatch.setattr("app.main.classify_site", fake_classify)
    client.post(f"/classify/{site_id}", follow_redirects=True)

    assert called["id"] == site_id
    with get_session() as session:
        assert session.get(Website, site_id).company_name == "ACME"


def test_exports_follow_the_filter_and_stream_rows(make_website, client):
    make_website(
        "https://acme.test",
        status="classified",
        company_name="ACME Robotics",
        industry="Manufacturing / Industrial",
        description="Builds warehouse robots.",
        confidence=88,
        clean_text="ACME builds warehouse robots.",
    )
    make_website("https://school.test", status="classified", company_name="School Co", industry="Education")

    csv_body = client.get("/export.csv", params={"q": "ACME"}).text
    assert csv_body.splitlines()[0].startswith("original_url,")
    assert "ACME Robotics" in csv_body
    assert "School Co" not in csv_body

    payload = json.loads(client.get("/export.json", params={"industry": "Education"}).text)
    assert len(payload) == 1
    assert payload[0]["company_name"] == "School Co"
    assert payload[0]["classification_status"] == "classified"
    assert payload[0]["crawl_status"] == ""

    sheet = load_workbook(io.BytesIO(client.get("/export.xlsx").content), read_only=True).active
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header[0] == "original_url"
    assert header[2] == "company_name"
    names = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "ACME Robotics" in names
    assert "School Co" in names
